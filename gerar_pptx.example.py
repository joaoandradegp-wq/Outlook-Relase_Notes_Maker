import copy
import os
import re
import tempfile
from pptx import Presentation
from pptx.oxml.ns import qn
from PIL import Image

CROP_PAINEL_LEITURA = (775, 145, 1265, 460)

def recortar_print_leitura(caminho_entrada):

    with Image.open(caminho_entrada) as im:
        recortada = im.crop(CROP_PAINEL_LEITURA) if im.size == (1280, 720) else im.copy()
        fd, caminho_saida = tempfile.mkstemp(suffix=".png", prefix="recorte_")
        os.close(fd)
        recortada.save(caminho_saida)
        return caminho_saida



def duplicar_slide(prs, indice_origem):

    slide_origem = prs.slides[indice_origem]
    layout = slide_origem.slide_layout
    novo_slide = prs.slides.add_slide(layout)

    for shape in list(novo_slide.shapes):
        shape._element.getparent().remove(shape._element)

    mapa_rid = {}
    for rel in slide_origem.part.rels.values():
        if "image" in rel.reltype:
            image_part = rel.target_part
            novo_rid = novo_slide.part.relate_to(image_part, rel.reltype)
            mapa_rid[rel.rId] = novo_rid

    for shape in slide_origem.shapes:
        el_copia = copy.deepcopy(shape._element)
        for blip in el_copia.iter(qn("a:blip")):
            rid_antigo = blip.get(qn("r:embed"))
            if rid_antigo and rid_antigo in mapa_rid:
                blip.set(qn("r:embed"), mapa_rid[rid_antigo])
        novo_slide.shapes._spTree.append(el_copia)

    return novo_slide


def remover_slide(prs, indice):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[indice])


def achar_forma(slide, nome):
    for shape in slide.shapes:
        if shape.name == nome:
            return shape
    raise ValueError(f'Forma "{nome}" não encontrada no slide.')


def achar_formas_imagem_ordenadas_por_x(slide):
    imagens = [s for s in slide.shapes if s.shape_type == 13]
    return sorted(imagens, key=lambda s: s.left)



PADRAO_DEMANDA = re.compile(r"^\s*(?:user\s+story\s*)?(\d{4,})\s*[:\-]\s*(.*)$", re.IGNORECASE)


def normalizar_demanda(texto):
    m = PADRAO_DEMANDA.match(texto)
    if m:
        numero, resto = m.groups()
        return f"User Story {numero}: ", resto.strip()
    return "", texto.strip()


def clonar_paragrafo_2_runs(paragrafo_modelo_xml, prefixo, resto):
    p = copy.deepcopy(paragrafo_modelo_xml)
    runs = p.findall(qn("a:r"))

    run_azul = runs[0]
    run_azul.find(qn("a:t")).text = prefixo

    run_preto = runs[2] if len(runs) > 2 else runs[-1]
    run_preto.find(qn("a:t")).text = resto

    for r in runs:
        if r is not run_azul and r is not run_preto:
            p.remove(r)

    return p


def preencher_lista_demandas(shape_caixa_texto, demandas):
    txBody = shape_caixa_texto.text_frame._txBody
    paragrafo_modelo = copy.deepcopy(txBody.findall(qn("a:p"))[0])

    for p_antigo in txBody.findall(qn("a:p")):
        txBody.remove(p_antigo)

    for texto in demandas:
        prefixo, resto = normalizar_demanda(texto)
        novo_p = clonar_paragrafo_2_runs(paragrafo_modelo, prefixo, resto)
        txBody.append(novo_p)


def definir_titulo_demanda(shape_caixa_texto, texto_demanda):
    txBody = shape_caixa_texto.text_frame._txBody
    paragrafo_modelo = copy.deepcopy(txBody.findall(qn("a:p"))[0])
    prefixo, resto = normalizar_demanda(texto_demanda)
    novo_p = clonar_paragrafo_2_runs(paragrafo_modelo, prefixo, resto)
    p_antigo = txBody.findall(qn("a:p"))[0]
    txBody.replace(p_antigo, novo_p)



def substituir_imagem(shape_imagem, caminho_novo_arquivo):
    slide_part = shape_imagem.part
    image_part, rId = slide_part.get_or_add_image_part(caminho_novo_arquivo)
    shape_imagem._element.blipFill.blip.rEmbed = rId

    largura_caixa, altura_caixa = shape_imagem.width, shape_imagem.height
    with Image.open(caminho_novo_arquivo) as img:
        largura_img, altura_img = img.size

    escala = min(largura_caixa / largura_img, altura_caixa / altura_img)
    shape_imagem.width = int(largura_img * escala)
    shape_imagem.height = int(altura_img * escala)



def indice_atual_do_slide(prs, slide_obj):
    alvo = slide_obj._element
    for i, s in enumerate(prs.slides):
        if s._element is alvo:
            return i
    raise ValueError("slide não encontrado na apresentação")


def gerar_pptx(template_path, output_path, mes_ano, demandas, aprovacoes):
    prs = Presentation(template_path)

    slide_capa = prs.slides[0]
    slide_lista = prs.slides[1]
    slide_demanda_modelo = prs.slides[2]
    slide_aprovacao_modelo = prs.slides[3]

    caixa_titulo = achar_forma(slide_capa, "CaixaDeTexto 3")
    paragrafo_mes = caixa_titulo.text_frame.paragraphs[1]
    paragrafo_mes.runs[0].text = mes_ano

    caixa_lista = achar_forma(slide_lista, "CaixaDeTexto 2")
    if demandas:
        preencher_lista_demandas(caixa_lista, demandas)

    if demandas:
        indice_modelo_demanda = 2
        slides_demanda = [slide_demanda_modelo]
        for _ in demandas[1:]:
            slides_demanda.append(duplicar_slide(prs, indice_modelo_demanda))
        for slide, texto in zip(slides_demanda, demandas):
            caixa = achar_forma(slide, "CaixaDeTexto 5")
            definir_titulo_demanda(caixa, texto)

    if aprovacoes:
        indice_modelo_aprovacao = 3
        slides_aprovacao = [slide_aprovacao_modelo]
        for _ in aprovacoes[1:]:
            slides_aprovacao.append(duplicar_slide(prs, indice_modelo_aprovacao))
        for slide, par in zip(slides_aprovacao, aprovacoes):
            img_pedido, img_resposta = achar_formas_imagem_ordenadas_por_x(slide)
            substituir_imagem(img_pedido, recortar_print_leitura(par["pedido"]))
            substituir_imagem(img_resposta, recortar_print_leitura(par["resposta"]))

    if not demandas:
        remover_slide(prs, indice_atual_do_slide(prs, slide_demanda_modelo))
    if not aprovacoes:
        remover_slide(prs, indice_atual_do_slide(prs, slide_aprovacao_modelo))

    prs.save(output_path)
    return output_path


if __name__ == "__main__":
    import json
    import sys

    if len(sys.argv) != 2:
        print("Uso: python3 gerar_pptx.py caminho/para/config.json", file=sys.stderr)
        sys.exit(1)

    with open(sys.argv[1], "r", encoding="utf-8") as f:
        cfg = json.load(f)

    caminho_gerado = gerar_pptx(
        template_path=cfg["template"],
        output_path=cfg["output"],
        mes_ano=cfg["mesAno"],
        demandas=cfg.get("demandas", []),
        aprovacoes=cfg.get("aprovacoes", []),
    )
    print(caminho_gerado)

